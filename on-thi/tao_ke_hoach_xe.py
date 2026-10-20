"""Tạo file Word "Kế hoạch sử dụng xe đi công tác": 2 phiếu khổ A5 đứng trên 1 tờ A4 ngang.

    pip install python-docx
    python tao_ke_hoach_xe.py [ke_hoach_xe.docx]

In A4 ngang, cắt đôi theo đường nét đứt ở giữa được 2 phiếu A5.
Các mục in sẵn giống mẫu; các mục còn lại để chấm chấm cho bạn ghi tay.
"""
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

FONT = "Times New Roman"
A5_W, A5_H = 14.8, 20.7          # cm (A5 = 14,8 x 21; trừ chút để không tràn sang trang 2)
PAD = 1.0                          # lề trong mỗi phiếu (cm)
CONTENT_W = A5_W - 2 * PAD

# (nhãn, giá trị in sẵn, số dòng chấm). Giá trị rỗng = để ghi tay.
FIELDS = [
    ("Đơn vị sử dụng: ", "Phòng Kế hoạch", 1),
    ("Cấp bậc, họ tên người sử dụng: ", "Hồ Nhật Vũ", 1),
    ("Chức vụ người sử dụng xe: ", "Phó Trưởng phòng", 1),
    ("Số người đi: ", "", 1),
    ("Số lượng hàng hóa vận chuyển: ", "", 1),
    ("Nội dung công tác: ", "", 1),
    ("Nơi đến công tác: ", "", 1),
    ("Loại xe: ", "07 chỗ", 1),
    ("Số KM đi và về: ", "", 1),
    ("Ngày đi: ", "", 1),
    ("Ngày về: ", "", 1),
    ("Địa điểm đón xe: ", "", 2),
    ("SĐT liên hệ: ", "0984.693.819", 1),
]


def set_font(run, size, bold=False, italic=False, underline=False):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold, run.italic, run.underline = bold, italic, underline


def para(container, text="", size=12, bold=False, italic=False, align=None,
         before=0, after=0, line=None, underline=False):
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    if line:
        pf.line_spacing = line
    if align is not None:
        p.alignment = align
    if text:
        set_font(p.add_run(text), size, bold, italic, underline)
    return p


def cell_border(cell, **edges):
    """edges: right='dashed' ... ; áp dụng đường viền cho ô."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge, style in edges.items():
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), style)
        el.set(qn("w:sz"), "6")
        el.set(qn("w:color"), "808080")
        borders.append(el)
    mar = tcPr.find(qn("w:tcMar"))
    if mar is not None:
        mar.addprevious(borders)
    else:
        tcPr.append(borders)


def cell_margins(cell, top=0, bottom=0, left=0, right=0):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{k}")
        el.set(qn("w:w"), str(int(v / 2.54 * 1440)))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)


def fix_widths(table, widths_cm):
    table.autofit = False
    for row in table.rows:
        for cell, w in zip(row.cells, widths_cm):
            cell.width = Cm(w)


def first_par(cell):
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    return p


def build_plan(cell):
    cell_margins(cell, top=0.9, bottom=0.6, left=PAD, right=PAD)

    # --- Quốc hiệu / tên cơ quan ---
    head = cell.add_table(rows=1, cols=2)
    fix_widths(head, [5.5, CONTENT_W - 5.5])
    left, right = head.rows[0].cells
    p = first_par(left); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("BỘ QUỐC PHÒNG"), 11)
    para(left, "VIỆN CHIẾN LƯỢC VÀ LỊCH SỬ", 11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(left, "QUỐC PHÒNG VIỆT NAM", 11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, underline=False)
    para(left, "Số: ............../KH-PKH.", 11, align=WD_ALIGN_PARAGRAPH.CENTER, before=2)
    p = first_par(right); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM"), 11, bold=True)
    para(right, "Độc lập - Tự do - Hạnh phúc", 12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, underline=True)
    para(right, "Hà Nội, ngày ....... tháng ....... năm 2026", 11, italic=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, before=10)

    # --- Tiêu đề ---
    para(cell, "KẾ HOẠCH", 14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=16)
    para(cell, "SỬ DỤNG XE ĐI CÔNG TÁC", 14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
         underline=True, after=8)

    # --- Các mục có dòng chấm ---
    for label, value, lines in FIELDS:
        p = para(cell, line=1.0, before=6)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(CONTENT_W), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        set_font(p.add_run(label), 12)
        if value:
            set_font(p.add_run(value), 12)
        set_font(p.add_run("\t"), 12)
        for _ in range(lines - 1):
            p = para(cell, line=1.0, before=6)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(CONTENT_W), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            set_font(p.add_run("\t"), 12)

    # --- Chữ ký ---
    para(cell, "", 6, before=6)
    sig = cell.add_table(rows=1, cols=2)
    fix_widths(sig, [CONTENT_W / 2 - 0.2, CONTENT_W / 2 + 0.2])
    a, b = sig.rows[0].cells
    p = first_par(a); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("NGƯỜI LẬP KẾ HOẠCH"), 12, bold=True)
    p = first_par(b); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("KT. TRƯỞNG PHÒNG"), 12, bold=True)
    para(b, "PHÓ TRƯỞNG PHÒNG", 12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(b, "", 12, before=0, after=0).paragraph_format.space_after = Pt(42)
    para(b, "Đại tá Hồ Nhật Vũ", 12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)


def main():
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).with_name("ke_hoach_xe.docx")
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = FONT, Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    sec.top_margin = sec.bottom_margin = Cm(0.05)
    sec.left_margin = sec.right_margin = Cm(0.05)

    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    fix_widths(t, [A5_W, A5_W])
    row = t.rows[0]
    row.height, row.height_rule = Cm(A5_H), WD_ROW_HEIGHT_RULE.EXACTLY
    for cell in row.cells:
        build_plan(cell)
    cell_border(row.cells[0], right="dashed")      # đường cắt giữa 2 phiếu

    # đoạn rỗng cuối bắt buộc của Word: thu nhỏ để không sinh trang trắng
    last = doc.add_paragraph()
    last.paragraph_format.space_before = last.paragraph_format.space_after = Pt(0)
    last.paragraph_format.line_spacing = Pt(1)
    set_font(last.add_run(" "), 1)

    doc.save(out)
    print("Đã tạo", out)


if __name__ == "__main__":
    main()
